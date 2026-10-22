"""Проверки обязательных дополнений к упрощённой версии."""
from io import BytesIO
from docx import Document
import pytest
from app.models import Attachment, Olympiad, Participation, User
from tests.test_olympiads import payload
from tests.test_attachments import upload


@pytest.mark.parametrize('version', ['v1', 'v2'])
def test_replace_olympiad(client, headers, accounts, version):
    base = f'/api/{version}/olympiads'
    data = payload(accounts['teacher'].id)
    event = client.post(base, json=data, headers=headers['admin']).json()
    data['description'] = 'Обновлённое описание'
    response = client.put(base + f'/{event["id"]}', json=data, headers=headers['admin'])
    assert response.status_code == 200
    assert response.json()['description'] == 'Обновлённое описание'
    assert client.put(base + f'/{event["id"]}', json=data, headers=headers['teacher']).status_code == 403


def test_edit_award_blank(client, db, headers, olympiad, tmp_path):
    response = upload(client, headers, olympiad, kind='award')
    fid = response.json()['id']
    old_name = db.get(Attachment, fid).stored_name
    stream = BytesIO()
    doc = Document()
    doc.add_paragraph('Новый бланк грамоты')
    doc.save(stream)
    url = f'/api/v1/attachments/{fid}'
    response = client.put(url, headers=headers['teacher'], data={'kind':'award','title':'Грамота'},
                          files={'file':('award.docx', stream.getvalue())})
    assert response.status_code == 200
    assert response.json()['id'] == fid
    assert not (tmp_path / old_name).exists()
    downloaded = client.get(url)
    assert Document(BytesIO(downloaded.content)).paragraphs[0].text == 'Новый бланк грамоты'
    assert client.put(url, headers=headers['other_teacher'], data={'kind':'award','title':'Чужая'},
                      files={'file':('test.txt', b'Test')}).status_code == 403


def test_demo_has_required_records(db, tmp_path, monkeypatch):
    import init_db
    monkeypatch.setattr(init_db, 'STORAGE', tmp_path)
    init_db.seed_demo(db)
    for model in (User, Olympiad, Participation, Attachment):
        assert db.query(model).count() >= 50
    assert db.query(User).filter_by(role='admin').count() == 1
    assert db.query(User).filter_by(role='teacher').count() >= 2
    kinds = {item.kind for item in db.query(Attachment).all()}
    assert kinds == {'regulation','task','certificate','award'}
    for item in db.query(Attachment).all():
        document = Document(tmp_path / item.stored_name)
        assert document.paragraphs[0].text == item.title
    before = [db.query(model).count() for model in (User, Olympiad, Participation, Attachment)]
    init_db.seed_demo(db)
    assert before == [db.query(model).count() for model in (User, Olympiad, Participation, Attachment)]
