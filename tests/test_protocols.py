from datetime import date, timedelta
from io import BytesIO
import pytest
from docx import Document
from openpyxl import load_workbook
from app.models.participation import Participation
from app.models.olympiad import Olympiad

BASE='/api/v1'


def enroll(client, headers, olympiad):
    response=client.post(BASE+f'/olympiads/{olympiad.id}/participations',headers=headers['student'])
    assert response.status_code==201
    return response.json()['id']


def score(client, headers, pid):
    return client.put(BASE+f'/participations/{pid}/result',json={'score':95.5,'result':'winner'},headers=headers['teacher'])


def test_complete_workflow(client, db, headers, accounts, olympiad):
    oid=olympiad.id; url=BASE+f'/olympiads/{oid}/protocol'
    pid=enroll(client,headers,olympiad)
    assert client.post(BASE+f'/olympiads/{oid}/participations',headers=headers['student']).status_code==409
    assert client.patch(url,json={'status':'prepared'},headers=headers['teacher']).status_code==409
    assert score(client,headers,pid).status_code==200
    assert client.get(BASE+'/users/me/participations',headers=headers['student']).json()[0]['score'] is None
    assert client.get(url).status_code==404
    assert client.get(BASE+f'/olympiads/{oid}/protocol-draft',headers=headers['teacher']).status_code==200
    assert client.patch(url,json={'status':'prepared'},headers=headers['teacher']).status_code==200
    assert score(client,headers,pid).status_code==409
    assert client.delete(BASE+f'/participations/{pid}',headers=headers['student']).status_code==409
    assert client.patch(url,json={'status':'draft'},headers=headers['teacher']).status_code==200
    assert score(client,headers,pid).status_code==200
    assert client.patch(url,json={'status':'published'},headers=headers['admin']).status_code==409
    assert client.patch(url,json={'status':'prepared'},headers=headers['teacher']).status_code==200
    assert client.patch(url,json={'status':'published'},headers=headers['teacher']).status_code==403
    assert client.patch(url,json={'status':'published'},headers=headers['admin']).status_code==200
    data=client.get(url).json(); assert data['participants'][0]['score']==95.5
    assert data['participants'][0]['result']=='winner'
    assert 'email' not in str(data) and 'birthDate' not in str(data)
    assert client.get(BASE+'/users/me/participations',headers=headers['student']).json()[0]['result']=='winner'
    assert client.patch(url,json={'status':'draft'},headers=headers['teacher']).status_code==409
    for format in ('docx','xlsx'):
        response=client.get(BASE+f'/olympiads/{oid}/reports/protocol',params={'format':format},headers=headers['teacher'])
        assert response.status_code==200
        if format=='docx':
            doc=Document(BytesIO(response.content)); assert doc.tables[0].rows[1].cells[1].text==accounts['student'].full_name
        else:
            book=load_workbook(BytesIO(response.content)); assert book.active.cell(3,3).value==95.5
        annual=client.get(BASE+'/reports/annual',params={'year':date.today().year,'format':format},headers=headers['admin'])
        assert annual.status_code==200
        if format=='xlsx':
            book=load_workbook(BytesIO(annual.content)); assert book.active.max_row==3
    assert client.get(BASE+'/reports/annual',params={'year':2000},headers=headers['admin']).status_code==200
    assert client.get(BASE+f'/olympiads/{oid}/reports/protocol',params={'format':'pdf'},headers=headers['admin']).status_code==422
    assert client.get(BASE+'/reports/annual',params={'year':date.today().year},headers=headers['teacher']).status_code==403


def test_ownership(client, headers, olympiad):
    pid=enroll(client,headers,olympiad); oid=olympiad.id
    other=headers['other_teacher']
    assert client.get(BASE+f'/olympiads/{oid}/participations',headers=other).status_code==403
    assert client.get(BASE+f'/olympiads/{oid}/protocol-draft',headers=other).status_code==403
    assert client.patch(BASE+f'/olympiads/{oid}/protocol',json={'status':'prepared'},headers=other).status_code==403
    assert client.put(BASE+f'/participations/{pid}/result',json={'score':10,'result':'participant'},headers=other).status_code==403
    assert client.delete(BASE+f'/participations/{pid}',headers=headers['other_student']).status_code==403
    assert client.get(BASE+f'/olympiads/{oid}/reports/protocol',headers=other).status_code==403
    assert client.get(BASE+f'/olympiads/{oid}/participations',headers=headers['student']).status_code==403
    assert len(client.get(BASE+f'/olympiads/{oid}/participations',headers=headers['teacher']).json())==1
    assert len(client.get(BASE+'/protocols',headers=headers['teacher']).json())==1
    assert client.get(BASE+'/protocols',headers=other).json()==[]
    assert len(client.get(BASE+'/protocols',headers=headers['admin']).json())==1


def test_enrollment_and_deletion(client,db,headers,olympiad):
    oid=olympiad.id
    assert client.post(BASE+f'/olympiads/{oid}/participations',headers=headers['teacher']).status_code==403
    pid=enroll(client,headers,olympiad)
    assert client.delete(BASE+f'/olympiads/{oid}',headers=headers['admin']).status_code==409
    assert client.delete(BASE+f'/participations/{pid}',headers=headers['student']).status_code==204
    assert client.get(BASE+'/users/me/participations',headers=headers['student']).json()==[]
    pid=enroll(client,headers,olympiad)
    assert client.delete(BASE+f'/participations/{pid}',headers=headers['teacher']).status_code==204
    olympiad.start_date=date.today()-timedelta(days=2); olympiad.end_date=date.today()-timedelta(days=1); db.commit()
    assert client.post(BASE+f'/olympiads/{oid}/participations',headers=headers['student']).status_code==409


def test_protocol_incomplete_future_and_validation(client,db,headers,olympiad):
    url=BASE+f'/olympiads/{olympiad.id}/protocol'
    assert client.patch(url,json={'status':'prepared'},headers=headers['teacher']).status_code==409
    pid=enroll(client,headers,olympiad)
    for data in [{'score':-1,'result':'winner'},{'score':10,'result':'unknown'},
                 {'score':1000000,'result':'winner'},{'score':1.111,'result':'winner'}]:
        assert client.put(BASE+f'/participations/{pid}/result',json=data,headers=headers['teacher']).status_code==422
    score(client,headers,pid)
    olympiad.end_date=date.today()+timedelta(days=1); db.commit()
    assert client.patch(url,json={'status':'prepared'},headers=headers['teacher']).status_code==409
    assert client.patch(url,json={'status':'invalid'},headers=headers['teacher']).status_code==422


@pytest.mark.parametrize('method,path,data', [('get','/olympiads/9999999',None),
    ('get','/olympiads/9999999/protocol',None), ('delete','/participations/9999999',None),
    ('put','/participations/9999999/result',{'score':10,'result':'winner'})])
def test_not_found(client,headers,method,path,data):
    kwargs={'headers':headers['teacher']}
    if data: kwargs['json']=data
    assert getattr(client,method)(BASE+path,**kwargs).status_code==404


def test_excel_formula_safety(client,db,headers,accounts,olympiad):
    accounts['student'].full_name='=HYPERLINK("bad")'; db.commit()
    pid=enroll(client,headers,olympiad); score(client,headers,pid)
    response=client.get(BASE+f'/olympiads/{olympiad.id}/reports/protocol?format=xlsx',headers=headers['teacher'])
    book=load_workbook(BytesIO(response.content))
    assert book.active.cell(3,2).data_type=='s'
    assert book.active.cell(3,2).value.startswith("'=")


def test_draft_export(client,headers,olympiad):
    enroll(client,headers,olympiad)
    response=client.get(BASE+f'/olympiads/{olympiad.id}/reports/protocol',headers=headers['teacher'])
    assert Document(BytesIO(response.content)).tables[0].rows[1].cells[2].text=='—'


def test_v2_security_and_write(client,headers,accounts):
    from tests.test_olympiads import payload
    response=client.post('/api/v2/olympiads',json=payload(accounts['teacher'].id),headers=headers['admin'])
    assert response.status_code==201
    oid=response.json()['id']
    assert client.get(f'/api/v2/olympiads/{oid}').status_code==200
    assert client.get('/api/v2/users',headers=headers['student']).status_code==403
    assert client.delete(f'/api/v2/olympiads/{oid}',headers=headers['admin']).status_code==204
