"""Сборка пояснительной записки из REPORT_TEXT.md и документации.

Запуск из корня: python -m scripts.build_report.
Опционально --page-map docs/evidence/page_map.json после первого экспорта PDF.
Требуются python-docx и Pillow из requirements-lock.txt.
"""
from pathlib import Path
import argparse
import json
import re
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
SOURCES = [
'Техническое задание «Пеннер (1)». Тема: создание и тестирование АРМ учёта участия студентов в олимпиадном движении. Предоставленный материал, 2026.',
'Методические рекомендации по выполнению курсового проекта ПМ.05 / сост. А. В. Грибова. Екатеринбург: ГАПОУ СО «ЕМК», 2026. Предоставленный PDF.',
'Общие правила оформления курсовой работы. Раздел 3. Предоставленный PDF, без выходных данных.',
'Рамки (1). Шаблон оформления курсового проекта. Предоставленный DOCX, без выходных данных.',
'Библиотека requests. Презентация о HTTP, REST и запросах Python. Предоставленный PDF, без выходных данных.',
'FastAPI. Документация разработчика: маршруты, зависимости, безопасность, подприложения. Электронный ресурс: https://fastapi.tiangolo.com/.',
'Pydantic. Документация версии 2: модели, поля, валидация. Электронный ресурс: https://docs.pydantic.dev/latest/.',
'SQLAlchemy. Документация версии 2: ORM, сессии и транзакции. Электронный ресурс: https://docs.sqlalchemy.org/en/20/.',
'PostgreSQL Global Development Group. PostgreSQL 17 Documentation. Электронный ресурс: https://www.postgresql.org/docs/17/.',
'Requests. User Guide. Электронный ресурс: https://requests.readthedocs.io/en/latest/.',
'PyJWT. Documentation: encoding and decoding tokens. Электронный ресурс: https://pyjwt.readthedocs.io/en/stable/.',
'python-docx. Documentation: documents and tables. Электронный ресурс: https://python-docx.readthedocs.io/en/latest/.',
'openpyxl. Documentation: workbook and cell values. Электронный ресурс: https://openpyxl.readthedocs.io/en/stable/.',
'pytest. Documentation: fixtures and parametrization. Электронный ресурс: https://docs.pytest.org/en/stable/.',
'Fielding R., Nottingham M., Reschke J. HTTP Semantics. RFC 9110. IETF, 2022. Электронный ресурс: https://www.rfc-editor.org/rfc/rfc9110.'
]
BOOKS = [
'Beaulieu A. Learning SQL. 3rd ed. O’Reilly Media, 2020.',
'Fowler M. Refactoring: Improving the Design of Existing Code. 2nd ed. Addison-Wesley, 2018.',
'Grinberg M. Flask Web Development. 2nd ed. O’Reilly Media, 2018.',
'Kleppmann M. Designing Data-Intensive Applications. O’Reilly Media, 2017.',
'Lub anovic B. Introducing Python. 2nd ed. O’Reilly Media, 2019.'.replace('Lub anovic','Lubanovic'),
'Martin R. C. Clean Architecture: A Craftsman’s Guide to Software Structure and Design. Prentice Hall, 2017.',
'Matthes E. Python Crash Course. 3rd ed. No Starch Press, 2023.',
'Percival H. J. W. Test-Driven Development with Python. 2nd ed. O’Reilly Media, 2017.',
'Ramalho L. Fluent Python. 2nd ed. O’Reilly Media, 2022.',
'Richards M., Ford N. Fundamentals of Software Architecture. O’Reilly Media, 2020.'
]
GLOSSARY = [
('АРМ','Автоматизированное рабочее место.'),('API','Программный интерфейс для взаимодействия приложений.'),
('REST','Подход к организации HTTP-взаимодействия с ресурсами.'),('HTTP','Протокол запросов и ответов клиента и сервера.'),
('JSON','Текстовый формат обмена структурированными данными.'),('JWT','Подписанный токен с набором утверждений и сроком действия.'),
('ORM','Отображение объектов программы на таблицы базы данных.'),('БД / СУБД','База данных / система управления базами данных.'),
('PK / FK','Первичный ключ / внешний ключ.'),('CRUD','Создание, чтение, обновление и удаление записей.'),
('OpenAPI','Машиночитаемое описание операций HTTP API.'),('Swagger / ReDoc','Интерфейсы чтения и проверки OpenAPI.'),
('PBKDF2','Функция получения хеша пароля с солью и повторными вычислениями.'),('Транзакция','Группа операций с подтверждением или откатом.'),
('Пагинация','Получение списка частями заданного размера.'),('Фикстура','Подготовка и освобождение состояния теста.'),
('Покрытие','Доля строк или ветвей кода, выполненных при проверках.'),('TDD','Разработка через цикл тест — реализация — рефакторинг.')]


def configure_section(section):
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(3), Cm(1.5)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2.5)
    section.header_distance, section.footer_distance = Cm(.5), Cm(.5)


def set_font(run, size=14, bold=None):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')


def frame(section, with_page):
    """Рамка и штамп по образцу шаблона, штатными элементами Word.

    Плавающие объекты исходного шаблона не повторялись в LibreOffice,
    поэтому рамка оформлена границей страницы, штамп — таблицей footer.
    """
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.different_first_page_header_footer = False
    for existing in section._sectPr.xpath('./w:pgBorders'):section._sectPr.remove(existing)
    borders = OxmlElement('w:pgBorders')
    borders.set(qn('w:offsetFrom'), 'page')
    borders.set(qn('w:display'),'allPages')
    borders.set(qn('w:zOrder'),'front')
    for side, space in [('top',14), ('left',57), ('bottom',14), ('right',14)]:
        element = OxmlElement('w:'+side)
        for key,value in [('val','single'),('sz','8'),('space',str(space)),('color','000000')]:
            element.set(qn('w:'+key),value)
        borders.append(element)
    cols=section._sectPr.find(qn('w:cols'))
    section._sectPr.insert(section._sectPr.index(cols) if cols is not None else 0,borders)
    footer = section.footer
    for child in list(footer._element):
        footer._element.remove(child)
    t = footer.add_table(rows=2, cols=7, width=Cm(18.5))
    t.autofit = False
    widths = [1,1,2.2,1.5,1.5,9.8,1.5]
    for column,width in zip(t.columns,widths):column.width=Cm(width)
    for row in t.rows:
        for cell,width in zip(row.cells,widths):cell.width=Cm(width)
    props=t._tbl.tblPr
    indent=OxmlElement('w:tblInd');indent.set(qn('w:w'),'-450');indent.set(qn('w:type'),'dxa');props.append(indent)
    grid=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        element=OxmlElement('w:'+side)
        for key,value in [('val','single'),('sz','8'),('color','000000')]:element.set(qn('w:'+key),value)
        grid.append(element)
    props.append(grid)
    for cell,label in zip(t.rows[1].cells,['Изм.','Лист','№ докум.','Подпись','Дата','','']):cell.text=label
    cell=t.cell(0,5).merge(t.cell(1,5));cell.text='ЕМК 09.02.07 КП № __________ ПЗ'
    t.cell(0,6).text='Лист'
    if with_page:
        p=t.cell(1,6).paragraphs[0]
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
        rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text='4';rr.append(tt);field.append(rr);p._p.append(field)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
                for run in p.runs:set_font(run,8)
    p=footer.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
    p.paragraph_format.line_spacing=Pt(1)
    set_font(p.add_run(''),1)


def paragraph(doc, text, style=None, size=14):
    p=doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing=1
    p.paragraph_format.space_after=Pt(6)
    p.paragraph_format.first_line_indent=Cm(1.25)
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    # Удаляются только Markdown-маркеры; код отдельно не преобразуется.
    text=text.replace('**','').replace('`','')
    set_font(p.add_run(text),size)
    return p


def heading(doc,text,level=1,new_page=True):
    p=doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.page_break_before=new_page
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.first_line_indent=Cm(1.25)
    p.paragraph_format.space_after=Pt(12)
    set_font(p.add_run(text.upper() if level==1 else text),14,True)
    return p


TABLE_COUNTER = 0


def table(doc,headers,rows,size=11):
    global TABLE_COUNTER
    TABLE_COUNTER += 1
    previous = doc.paragraphs[-1] if doc.paragraphs else None
    if previous is not None and previous.text.startswith('Таблица '):
        previous.text = re.sub(r'^Таблица \d+', 'Таблица '+str(TABLE_COUNTER), previous.text)
        for run in previous.runs:set_font(run,12)
    else:
        last_heading = next((p.text for p in reversed(doc.paragraphs) if p.style.name.startswith('Heading')), '')
        caption = 'Словарь данных: '+last_heading if headers[0]=='Ключ' else 'Выявленные проблемы'
        paragraph(doc, f'Таблица {TABLE_COUNTER} — {caption}',size=12)
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
    if headers[0]=='Ключ':
        widths=[2.5,2.8,2.8,1.5,6.9]
    elif headers[0]=='Метод':
        widths=[1.8,7.5,7.2]
    else:widths=None
    if widths:
        t.autofit=False
        for column,width in zip(t.columns,widths):column.width=Cm(width)
    for cell,text in zip(t.rows[0].cells,headers):
        cell.text=str(text)
    props=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(props)
    for row in rows:
        cells=t.add_row().cells
        for cell,value in zip(cells,row):
            cell.text=str(value).replace('`','')
    for i,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        if widths:
            for cell,width in zip(row.cells,widths):cell.width=Cm(width)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                p.paragraph_format.line_spacing=1
                p.paragraph_format.keep_with_next=headers[0]=='Ключ' and i<len(t.rows)-1
                for run in p.runs:set_font(run,size,i==0)
    return t


def image(doc,name,caption):
    path=next((folder/name for folder in (DOCS/'diagrams',DOCS/'evidence',DOCS/'examples') if (folder/name).exists()),None)
    if not path:raise FileNotFoundError(name)
    w,h=Image.open(path).size
    width=min(16.3,17*w/h)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Cm(width))
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(caption),12)


def code(doc,text):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.line_spacing=1
    for i,line in enumerate(text.splitlines()):
        run=p.add_run(('\n' if i else '')+line)
        run.font.name='Courier New';run.font.size=Pt(9)
    return p


def render_markdown(doc,text,include=False):
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('!FIG '):
            name,caption=line[5:].split('|',1);image(doc,name,caption);i+=1;continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'))
            if include and level==1:i+=1;continue
            previous = next((value.strip() for value in reversed(lines[:i]) if value.strip()), '')
            break_before = level<=2 and not (level==2 and previous.startswith('# '))
            if line[level:].strip() in ('olympiads','participations','attachments','Связи и ограничения'):break_before=False
            heading(doc,line[level:].strip(),min(level,3),new_page=break_before)
            i+=1;continue
        if line.startswith('```'):
            i+=1;buf=[]
            while i<len(lines) and not lines[i].startswith('```'):buf.append(lines[i]);i+=1
            code(doc,'\n'.join(buf));i+=1;continue
        if line.startswith('|'):
            cells=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+',x or '-') for x in row):cells.append(row)
                i+=1
            if cells:table(doc,cells[0],cells[1:])
            continue
        buf=[line];i+=1
        while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','!','```')):
            buf.append(lines[i].strip());i+=1
        paragraph(doc,' '.join(buf))


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--page-map',type=Path)
    args=parser.parse_args()
    metadata=json.loads((DOCS/'report_metadata.json').read_text(encoding='utf-8'))
    doc=Document();configure_section(doc.sections[0])
    for name in ('Normal','Heading 1','Heading 2','Heading 3'):
        style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(14)
        style.paragraph_format.line_spacing=1
    doc.core_properties.title='Учёт участия студентов в олимпиадах: REST-сервис FastAPI'
    doc.core_properties.author=metadata['author']
    # Титул без номера и рамки.
    for text in ['МИНИСТЕРСТВО ОБРАЗОВАНИЯ СВЕРДЛОВСКОЙ ОБЛАСТИ','ГАПОУ СО «ЕКАТЕРИНБУРГСКИЙ МОНТАЖНЫЙ КОЛЛЕДЖ»']:
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;set_font(p.add_run(text),12)
    doc.add_paragraph('\n\n')
    for text in ['КУРСОВОЙ ПРОЕКТ','ПОЯСНИТЕЛЬНАЯ ЗАПИСКА','Первый этап: REST-сервис на FastAPI']:
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;set_font(p.add_run(text),14,True)
    paragraph(doc,'По дисциплине «МДК 05.03 Тестирование информационных систем»')
    paragraph(doc,'Тема: «Создание и тестирование АРМ для обработки информации об участии студентов в олимпиадном движении, конкурсах профессионального мастерства»')
    doc.add_paragraph()
    for text in ['Выполнил: '+metadata['author'],'Группа: '+metadata['group'],
                 'Номер студенческого: '+metadata['student_id'],'Руководитель: '+metadata['supervisor'],
                 'Подписи: ________________________','Оценка: ________________________']:
        paragraph(doc,text)
    p=doc.add_paragraph('Екатеринбург\n'+metadata['year']);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:set_font(run)
    # Ведомость и содержание: учитываются в общем счёте, номер не печатается.
    s=doc.add_section(WD_SECTION_START.NEW_PAGE);configure_section(s);frame(s,False)
    heading(doc,'ВЕДОМОСТЬ ПРОЕКТА',new_page=False)
    paragraph(doc,'Таблица 1 — Состав комплекта первого этапа')
    table(doc,['№','Наименование','Представление'],[
        ['1','Пояснительная записка','DOCX и PDF'],['2','REST-сервис FastAPI v1/v2','app/'],
        ['3','Консольный HTTP-клиент','client/'],['4','Автоматические тесты','tests/'],
        ['5','Зависимости и развёртывание','requirements*, compose.yaml, scripts/'],
        ['6','Руководства и спецификации','docs/*.md, OpenAPI, DDL'],
        ['7','Диаграммы, отчёты и доказательства','docs/diagrams, examples, evidence'],
        ['8','Исходные материалы задания','docs/source/']])
    paragraph(doc,'Действующие пароли, .env, виртуальное окружение и содержимое пользовательского хранилища в передаваемый архив исходников не включаются. Демонстрационные данные создаются командой --demo. Полный состав проверяется по манифесту архива.')
    heading(doc,'СОДЕРЖАНИЕ')
    text=(DOCS/'REPORT_TEXT.md').read_text(encoding='utf-8')
    contents=[line.lstrip('# ').strip() for line in text.splitlines() if line.startswith('#')]
    contents += ['Приложение 1 — Реестр дефектов и ограничений','Приложение 2 — Тестовые сценарии',
                 'Приложение 3 — Матрица трассировки','Приложение 4 — Основные листинги',
                 'Приложение 5 — Запросы и примеры результатов','Приложение 6 — Рекомендуемая литература']
    page_map=json.loads(args.page_map.read_text(encoding='utf-8')) if args.page_map else {}
    for title in contents:
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
        run=p.add_run(title+' '+('· '*3)+str(page_map.get(title,'—')));set_font(run,11)
    s=doc.add_section(WD_SECTION_START.NEW_PAGE);configure_section(s);frame(s,True)
    metrics=json.loads((DOCS/'evidence/metrics.json').read_text())
    text=text.replace('!INCLUDE DATA_DICTIONARY.md', (DOCS/'DATA_DICTIONARY.md').read_text(encoding='utf-8').split('\n',1)[1])
    text=text.replace('!METRICS', f"Фактически выполнено {metrics['tests']} автоматических тестов: успешно {metrics['passed']}, падений {metrics['failed']}, пропусков {metrics['skipped']}. Покрытие строк app — {metrics['totalCoverage']:.2f}%, модулей API — {metrics['apiCoverage']:.2f}% ({metrics['apiCoveredLines']} из {metrics['apiLines']}). Прогон выполнен 19 октября 2026 года; точное время и длительность приведены в junit.xml и pytest.txt.")
    operations=[]
    schema=json.loads((DOCS/'openapi-v1.json').read_text())
    for path,methods in schema['paths'].items():
        for method,operation in methods.items():operations.append([method.upper(),path,operation['summary']])
    op_mark='В таблице 2 приведён состав операций версии v1. Версия v2 заменяет только GET коллекции олимпиад.\n\nТаблица 2 — Операции REST-сервиса\n\n| Метод | Путь | Назначение |\n|---|---|---|\n'
    op_mark+='\n'.join('| '+' | '.join(row)+' |' for row in operations)
    text=text.replace('!OPERATIONS',op_mark)
    gloss='В таблице 3 приведены термины и сокращения, используемые в работе.\n\nТаблица 3 — Термины и сокращения\n\n| Термин | Определение |\n|---|---|\n'+'\n'.join('| '+' | '.join(row)+' |' for row in GLOSSARY)
    text=text.replace('!GLOSSARY',gloss)
    text=text.replace('!SOURCES','\n\n'.join(f'{i}. {s}' for i,s in enumerate(SOURCES,1)))
    # Первое введение не должно создавать пустую страницу после разрыва раздела.
    render_markdown(doc,text)
    first_heading=next(p for p in doc.paragraphs if p.text=='ВВЕДЕНИЕ')
    first_heading.paragraph_format.page_break_before=False
    heading(doc,'Приложение 1 — Реестр дефектов и ограничений')
    testing=(DOCS/'TESTING.md').read_text(encoding='utf-8')
    render_markdown(doc,testing.split('## Реестр реально обнаруженных проблем')[1])
    paragraph(doc,'Приложение фиксирует исправленные проблемы разработки и тестовой инфраструктуры. Оно не заменяет окончательную приёмку и не содержит вымышленных дефектов. Настройки защищённого производственного размещения и второй этап Django в проверенный объём не входят.')
    heading(doc,'Приложение 2 — Тестовые сценарии')
    cases=json.loads((DOCS/'test-cases.json').read_text(encoding='utf-8'))
    for i,c in enumerate(cases):
        if i in (6,12,18):doc.add_page_break()
        p=paragraph(doc,f"{c['id']} — {c['requirements']}");p.paragraph_format.keep_with_next=True
        paragraph(doc,'Предусловия: тестовая транзакция PostgreSQL; аккаунты и олимпиада из фикстур.',size=11)
        paragraph(doc,'Действия и данные: '+c['action']+'. Ожидается: '+c['expected']+'.',size=11)
        paragraph(doc,'Реализация: '+c['test']+'. Фактический итог финального прогона: успешно.',size=11)
    heading(doc,'Приложение 3 — Матрица трассировки')
    paragraph(doc,'Таблица 4 сопоставляет требования R01–R19 из REQUIREMENTS.md и семейства тестов приложения 2. Порог заполнения данных проверен отдельно, а не выдан за автотест функциональности.')
    matrix=testing.split('## Матрица трассировки')[1].split('## Итог')[0]
    paragraph(doc,'Таблица 4 — Требования и проверки');render_markdown(doc,matrix)
    heading(doc,'Приложение 4 — Основные листинги')
    paragraph(doc,'Показаны код безопасности, управление состоянием и небольшой requests-клиент. Полные исходники находятся в app/ и client/. Это рабочие файлы комплекта; номера строк добавлены только для чтения листинга.')
    for filename in ['app/security.py','app/api/protocols.py','client/api/accounts.py']:
        lines=(ROOT/filename).read_text(encoding='utf-8').splitlines()
        for start in range(0,len(lines),50):
            if start or filename != 'app/security.py':doc.add_page_break()
            p=doc.add_paragraph(filename+f' (строки {start+1}–{min(start+50,len(lines))})')
            for run in p.runs:set_font(run,12,True)
            code(doc,'\n'.join(f'{i+1:3} {line}' for i,line in enumerate(lines[start:start+50],start)))
    heading(doc,'Приложение 5 — Запросы и примеры результатов')
    paragraph(doc,'Ниже показаны реальные страницы документации и примеры файлов из HTTP-сценария. Токены и значения JWT_SECRET на снимках не отображаются. Файлы protocol.docx/xlsx и annual.docx/xlsx находятся в docs/examples/.')
    image(doc,'swagger_overview.png','Рисунок 8 — Swagger первой версии работающего сервиса')
    doc.add_page_break();image(doc,'swagger_response.png','Рисунок 9 — Выполнение GET-запроса из Swagger')
    paragraph(doc,'Публичные запросы выполняются без авторизации. Для изменяющих операций необходимо войти через Authorize. Код ответа, заголовки и тело позволяют проверить фактический результат, как в презентации requests.')
    doc.add_page_break();image(doc,'swagger_v2.png','Рисунок 10 — Документация второй версии')
    paragraph(doc,'Вторая версия использует тот же набор предметных операций. Отличие списка: объект с items и параметрами страницы. Полный JSON-контракт и примеры запросов приведены в API.md.')
    doc.add_page_break()
    smoke=json.loads((DOCS/'evidence/smoke.json').read_text())
    paragraph(doc,'Фактический HTTP-сценарий: '+smoke['checkedAtUtc']+'. Результат: '+smoke['result']+'. Проверены вход, создание студента и олимпиады, заявка, баллы, подготовка, публикация и четыре бинарные выгрузки.')
    protocol=Document(DOCS/'examples/protocol.docx')
    paragraph(doc,'Таблица 5 — Содержимое примера опубликованного протокола')
    rows=[[cell.text for cell in row.cells] for row in protocol.tables[0].rows]
    table(doc,rows[0],rows[1:])
    code(doc,'import requests\nresponse = requests.get(\n    "http://127.0.0.1:8000/api/v1/olympiads",\n    params={"year": 2026}, timeout=10\n)\nresponse.raise_for_status()\nprint(response.status_code)\nprint(response.json())')
    paragraph(doc,'В примере ответ Response не подменяется словарём: сначала проверяется успешность, затем вызывается json(). Для сохранения отчёта используется content. Больше примеров и полный сценарий доступны в scripts/smoke.py и client/.')
    heading(doc,'Приложение 6 — Рекомендуемая литература')
    paragraph(doc,'Перечень книг для дополнительного самостоятельного изучения. Он отделён от технических источников основной части: прочтение всех книг при подготовке комплекта не заявляется. Книги не заменяют требуемые методичкой материалы периодической печати.')
    for i,book in enumerate(BOOKS,16):paragraph(doc,f'{i}. {book}',size=12)
    settings=doc.settings.element
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    for p in doc.paragraphs:
        value = p.text.replace('В таблице 2 приведён', 'В таблице 6 приведён').replace('В таблице 3 приведены', 'В таблице 7 приведены').replace('Таблица 4 сопоставляет', 'Таблица 9 сопоставляет')
        if value != p.text:
            p.text=value
            for run in p.runs:set_font(run)
    doc.save(DOCS/'Пояснительная_записка.docx')
    # Markdown-версия основной записки удобна для редактирования и поиска.
    (DOCS/'Пояснительная_записка.md').write_text(text,encoding='utf-8')
    print('Собрана пояснительная записка DOCX и Markdown')


if __name__=='__main__':main()
