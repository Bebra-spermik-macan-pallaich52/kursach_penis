"""Создать таблицы и учебные данные: python init_db.py. Непустая база не меняется."""
from datetime import date
from uuid import uuid4
from docx import Document
from app.database import Base, STORAGE, engine, SessionLocal
from app.models import Attachment, Olympiad, Participation, User
from app.auth import hash_password


def seed_demo(db):
    if db.query(User).first():
        print("В базе уже есть пользователи. Данные не изменены.")
        return
    admin = User(username="admin", password_hash=hash_password("Admin2026!"),
                 full_name="Администратор системы", email="admin@example.org", role="admin")
    db.add(admin)
    db.flush()
    teachers = []
    for i in range(1, 3):
        teacher = User(username=f"teacher{i}", password_hash=hash_password("Teacher2026!"),
                       full_name=f"Преподаватель Учебный Номер {i}", email=f"teacher{i}@example.org",
                       role="teacher", created_by=admin.id)
        db.add(teacher)
        db.flush()
        teachers.append(teacher)
    students = []
    for i in range(1, 48):
        student = User(username=f"student{i}", password_hash=hash_password("Student2026!"),
                       full_name=f"Студент Учебный Номер {i}", email=f"student{i}@example.org", role="student",
                       birth_date=date(2007, 1 + i % 12, 1 + i % 28), institution="Учебный колледж",
                       education_level="Среднее профессиональное", course=1 + i % 4,
                       specialty="09.02.07 Информационные системы и программирование",
                       created_by=teachers[i % 2].id)
        db.add(student)
        db.flush()
        students.append(student)
    today = date.today()
    kinds = ["regulation", "task", "certificate", "award"]
    titles = ["Положение", "Задание", "Сертификат", "Грамота"]
    STORAGE.mkdir(parents=True, exist_ok=True)
    created_files = []
    try:
        for i in range(50):
            event_date = today if i == 0 else date(today.year - 1, 1 + i % 12, 1 + i % 28)
            event = Olympiad(title=f"Олимпиада по программированию {i + 1}",
                             description="Вымышленное учебное мероприятие", teacher_id=teachers[i % 2].id,
                             start_date=event_date, end_date=event_date,
                             protocol_status=["draft", "published", "prepared"][i % 3])
            db.add(event)
            db.flush()
            db.add(Participation(olympiad_id=event.id, student_id=students[(i + 1) % 47].id,
                                 score=70 + i % 30, result=["winner", "prize", "participant"][i % 3]))
            document = Document()
            document.add_heading(titles[i % 4], 0)
            document.add_paragraph(event.title)
            if i % 4 == 0:
                document.add_paragraph("К участию допускаются студенты. Работы оценивает ответственный преподаватель. Итоги публикуются после проверки администратором.")
            elif i % 4 == 1:
                document.add_paragraph("Написать программу подсчёта баллов участников. Максимальная оценка — 100 баллов.")
            else:
                document.add_paragraph("Награждается: __________________. Дата и подпись: __________________.")
            name = uuid4().hex + ".docx"
            path = STORAGE / name
            created_files.append(path)
            document.save(path)
            db.add(Attachment(olympiad_id=event.id, kind=kinds[i % 4], title=titles[i % 4],
                              original_name=kinds[i % 4] + ".docx", stored_name=name))
        db.commit()
    except Exception:
        db.rollback()
        for path in created_files:
            path.unlink(missing_ok=True)
        raise
    print("Созданы 50 пользователей, 50 олимпиад, 50 заявок, 50 документов.")
    print("Учебные аккаунты и пароли приведены в README.md.")


def main():
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        seed_demo(db)


if __name__ == "__main__":
    main()
