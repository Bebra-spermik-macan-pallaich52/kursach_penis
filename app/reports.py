from io import BytesIO
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font
from fastapi import HTTPException
from fastapi.responses import Response

RESULT_NAMES = {"winner": "Победитель", "prize": "Призёр", "participant": "Участник"}


def make_report(title, headers, rows, file_format, filename):
    stream = BytesIO()
    if file_format == "docx":
        document = Document()
        document.styles["Normal"].font.name = "Times New Roman"
        document.styles["Normal"].font.size = Pt(12)
        document.add_heading(title, 0)
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
        for row in rows:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = str(value)
        document.save(stream)
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif file_format == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "Отчёт"
        sheet.append([title])
        sheet.append(headers)
        for row in rows:
            # Текст пользователя не должен выполняться как формула Excel.
            sheet.append(["'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value for value in row])
        for cell in sheet[2]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A3"
        sheet.auto_filter.ref = f"A2:{sheet.cell(sheet.max_row, len(headers)).coordinate}"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 25
        book.save(stream)
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        raise HTTPException(422, "Допустимые форматы: docx, xlsx")
    return Response(stream.getvalue(), media_type=mime,
                    headers={"Content-Disposition": f'attachment; filename="{filename}.{file_format}"'})

