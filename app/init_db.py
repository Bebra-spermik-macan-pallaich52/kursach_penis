"""Создание таблиц и учебных данных: python -m app.init_db --demo."""
from datetime import date
from io import BytesIO
import argparse
import os
from uuid import uuid4
from docx import Document
from app.config import STORAGE
from app.database import Base, engine, SessionLocal
from app.models.user import User
from app.models.olympiad import Olympiad
from app.models.participation import Participation
from app.models.attachment import Attachment
from app.security import hash_password


def add_user(db, username, password, role, full_name, **extra):
    user = User(username=username, password_hash=hash_password(password), role=role,
                full_name=full_name, email=username + "@example.org", **extra)
    db.add(user)
    db.flush()
    return user


def seed_demo(db):
    if db.query(User).first():
        print("База непустая: демонстрационные данные не добавлены и не изменены")
        return
    admin = add_user(db, "admin", "Admin2026!", "admin", "Администратор учебной системы")
    teachers = []
    for i in range(1, 51):
        teachers.append(add_user(db, f"teacher{i}", "Teacher2026!", "teacher",
                        f"Преподаватель Тестовый Номер {i}", created_by=admin.id))
    surnames = ["Иванов", "Петров", "Сидоров", "Смирнов", "Кузнецов", "Попов",
                "Васильев", "Соколов", "Михайлов", "Новиков"]
    students = []
    for i in range(1, 61):
        students.append(add_user(db, f"student{i}", "Student2026!", "student",
            f"{surnames[(i - 1) % 10]} Алексей Учебный {i}", birth_date=date(2006 + i % 3, 1 + i % 12, 1 + i % 25),
            institution="Екатеринбургский монтажный колледж (учебные данные)",
            education_level="Среднее профессиональное", course=1 + i % 4,
            specialty="09.02.07 Информационные системы и программирование", created_by=teachers[(i - 1) % 50].id))
    subjects = ["Программирование на Python", "Веб-разработка", "Базы данных", "Компьютерные сети", "Информационная безопасность"]
    STORAGE.mkdir(parents=True, exist_ok=True)
    created_files = []
    try:
        for i in range(50):
            event_date = date.today() if i == 0 else date(2024 + i % 3, 3 + i % 6, 1 + i % 25)
            olympiad = Olympiad(title=f"{subjects[i % 5]} — учебная олимпиада {i + 1}",
                description="Вымышленное мероприятие для демонстрации и тестирования.",
                teacher_id=teachers[i].id, start_date=event_date, end_date=event_date,
                protocol_status=["draft", "published", "prepared"][i % 3])
            db.add(olympiad)
            db.flush()
            for k in range(3):
                db.add(Participation(olympiad_id=olympiad.id,
                    student_id=students[(i * 3 + k) % 60].id,
                    score=95 - k * 15, result=["winner", "prize", "participant"][k]))
            for kind, title in [("regulation", "Положение"), ("task", "Задание"),
                                ("certificate", "Сертификат"), ("award", "Грамота")]:
                document = Document()
                document.add_heading(title, 0)
                document.add_paragraph(olympiad.title)
                document.add_paragraph("Учебный образец. Данные вымышлены.")
                if kind == "regulation":
                    document.add_paragraph("Цель: развитие профессиональных навыков. К участию допускаются студенты. Результат определяет преподаватель по утверждённым критериям.")
                elif kind == "task":
                    document.add_paragraph("Разработать программу обработки списка участников и подсчёта суммы баллов. Максимум — 100 баллов.")
                else:
                    document.add_paragraph("Награждается: __________________. Результат: __________. Дата и подпись: __________.")
                stream = BytesIO()
                document.save(stream)
                name = uuid4().hex + ".docx"
                path = STORAGE / name
                path.write_bytes(stream.getvalue())
                created_files.append(path)
                db.add(Attachment(olympiad_id=olympiad.id, kind=kind, title=title,
                    original_name=kind + ".docx", stored_name=name))
        db.commit()
    except Exception:
        db.rollback()
        for path in created_files:
            path.unlink(missing_ok=True)
        raise
    print("Создано: 111 пользователей, 50 олимпиад, 150 заявок, 200 документов")
    print("Учебные аккаунты: admin, teacher1..teacher50, student1..student60. Пароли — в README")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--demo", action="store_true", help="Заполнить только пустую БД учебными данными")
    args = parser.parse_args()
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        if args.demo:
            seed_demo(db)
        elif not db.query(User).first():
            username = os.getenv("ADMIN_USERNAME", "admin")
            password = os.getenv("ADMIN_PASSWORD", "")
            if len(password) < 8:
                raise RuntimeError("Укажите ADMIN_PASSWORD длиной от 8 символов в .env")
            add_user(db, username, password, "admin", "Администратор системы")
            db.commit()
            print("Таблицы и учётная запись администратора созданы")
        else:
            print("Таблицы уже существуют; пользовательские данные сохранены")


if __name__ == "__main__":
    main()

