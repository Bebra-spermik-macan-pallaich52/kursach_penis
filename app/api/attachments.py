from datetime import date
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile, BadZipFile
from docx import Document
from fastapi import APIRouter, Depends, File, Form, HTTPException, Response, UploadFile
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from app.config import MAX_FILE_SIZE, STORAGE
from app.database import get_db
from app.models.attachment import Attachment
from app.models.user import User
from app.schemas.attachment_read import AttachmentRead
from app.schemas.api_model import has_invalid_text
from app.security import require_staff
from app.services import check_owner, get_olympiad, save

router = APIRouter(tags=["Документы и бланки"])
KINDS = ("regulation", "task", "certificate", "award")
EXTENSIONS = (".pdf", ".docx", ".xlsx", ".png", ".jpg", ".jpeg", ".txt")


def validate_file(filename, content):
    extension = Path(filename).suffix.lower()
    if extension not in EXTENSIONS or not content:
        raise HTTPException(422, "Нужен непустой PDF, DOCX, XLSX, PNG, JPG или TXT")
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(413, "Размер файла превышает 5 МБ")
    valid = True
    if extension == ".pdf":
        valid = content.startswith(b"%PDF-")
    elif extension == ".png":
        valid = content.startswith(b"\x89PNG\r\n\x1a\n")
    elif extension in (".jpg", ".jpeg"):
        valid = content.startswith(b"\xff\xd8\xff")
    elif extension in (".docx", ".xlsx"):
        try:
            with ZipFile(BytesIO(content)) as archive:
                expected = "word/document.xml" if extension == ".docx" else "xl/workbook.xml"
                valid = expected in archive.namelist() and sum(i.file_size for i in archive.infolist()) <= 20 * 1024 * 1024
        except BadZipFile:
            valid = False
    elif extension == ".txt":
        try:
            content.decode("utf-8")
        except UnicodeDecodeError:
            valid = False
    if not valid:
        raise HTTPException(422, "Содержимое файла не соответствует расширению")
    return extension


def store_file(db, olympiad_id, kind, title, filename, content, item=None):
    if any(has_invalid_text(value) for value in (kind, title, filename)):
        raise HTTPException(422, "Текст содержит недопустимые управляющие символы")
    if kind not in KINDS or not 1 <= len(title.strip()) <= 150:
        raise HTTPException(422, "Проверьте тип и название документа")
    filename = filename.replace("\\", "/").split("/")[-1]
    if len(filename) > 200:
        raise HTTPException(422, "Слишком длинное имя файла")
    extension = validate_file(filename, content)
    STORAGE.mkdir(parents=True, exist_ok=True)
    stored_name = uuid4().hex + extension
    (STORAGE / stored_name).write_bytes(content)
    old_name = item.stored_name if item else None
    if item is None:
        item = Attachment(olympiad_id=olympiad_id)
        db.add(item)
    item.kind, item.title = kind, title.strip()
    item.original_name, item.stored_name = filename, stored_name
    try:
        save(db)
    except Exception:
        (STORAGE / stored_name).unlink(missing_ok=True)
        raise
    if old_name:
        (STORAGE / old_name).unlink(missing_ok=True)
    return item


def get_attachment(db, attachment_id):
    item = db.get(Attachment, attachment_id)
    if not item:
        raise HTTPException(404, "Документ не найден")
    return item


@router.get("/olympiads/{olympiad_id}/attachments", response_model=list[AttachmentRead], summary="Положение и бланки; задания прошлых лет")
def list_files(olympiad_id: int, db: Session = Depends(get_db)):
    olympiad = get_olympiad(db, olympiad_id)
    query = db.query(Attachment).filter_by(olympiad_id=olympiad.id)
    if olympiad.start_date.year >= date.today().year:
        query = query.filter(Attachment.kind != "task")
    return query.order_by(Attachment.id).all()


@router.get("/attachments/{attachment_id}", response_class=FileResponse,
    responses={200: {"content": {"application/octet-stream": {
        "schema": {"type": "string", "format": "binary"}}}}},
    summary="Скачать общедоступный документ")
def download(attachment_id: int, db: Session = Depends(get_db)):
    item = get_attachment(db, attachment_id)
    olympiad = get_olympiad(db, item.olympiad_id)
    if item.kind == "task" and olympiad.start_date.year >= date.today().year:
        raise HTTPException(404, "Доступны задания только прошлых лет")
    path = STORAGE / item.stored_name
    if not path.is_file():
        raise HTTPException(404, "Файл отсутствует в хранилище")
    return FileResponse(path, filename=item.original_name, media_type="application/octet-stream",
                         headers={"X-Content-Type-Options": "nosniff"})


@router.get("/olympiads/{olympiad_id}/attachments/manage", response_model=list[AttachmentRead], summary="Все файлы ответственного преподавателя")
def manage_files(olympiad_id: int, user: User = Depends(require_staff), db: Session = Depends(get_db)):
    olympiad = get_olympiad(db, olympiad_id)
    check_owner(olympiad, user)
    return db.query(Attachment).filter_by(olympiad_id=olympiad_id).order_by(Attachment.id).all()


@router.post("/olympiads/{olympiad_id}/attachments", response_model=AttachmentRead, status_code=201, summary="Прикрепить файл")
def upload(olympiad_id: int, kind: str = Form(), title: str = Form(), file: UploadFile = File(),
            user: User = Depends(require_staff), db: Session = Depends(get_db)):
    check_owner(get_olympiad(db, olympiad_id), user)
    content = file.file.read(MAX_FILE_SIZE + 1)
    return store_file(db, olympiad_id, kind, title, file.filename or "", content)


@router.put("/attachments/{attachment_id}", response_model=AttachmentRead, summary="Заменить файл или название")
def replace(attachment_id: int, kind: str = Form(), title: str = Form(), file: UploadFile = File(),
             user: User = Depends(require_staff), db: Session = Depends(get_db)):
    item = get_attachment(db, attachment_id)
    check_owner(get_olympiad(db, item.olympiad_id), user)
    return store_file(db, item.olympiad_id, kind, title, file.filename or "", file.file.read(MAX_FILE_SIZE + 1), item)


@router.post("/olympiads/{olympiad_id}/templates", response_model=AttachmentRead, status_code=201, summary="Создать положение или бланк Word из текста")
def create_template(olympiad_id: int, kind: str = Form(), title: str = Form(max_length=150),
                      text: str = Form(min_length=1, max_length=10000),
                      user: User = Depends(require_staff), db: Session = Depends(get_db)):
    olympiad = get_olympiad(db, olympiad_id)
    check_owner(olympiad, user)
    if any(has_invalid_text(value) for value in (kind, title, text)):
        raise HTTPException(422, "Текст содержит недопустимые управляющие символы")
    document = Document()
    document.add_heading(title, 0)
    document.add_paragraph(olympiad.title)
    document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    return store_file(db, olympiad_id, kind, title, "template.docx", stream.getvalue())


@router.delete("/attachments/{attachment_id}", status_code=204, summary="Удалить документ")
def delete_file(attachment_id: int, user: User = Depends(require_staff), db: Session = Depends(get_db)):
    item = get_attachment(db, attachment_id)
    check_owner(get_olympiad(db, item.olympiad_id), user)
    path = STORAGE / item.stored_name
    db.delete(item)
    save(db)
    path.unlink(missing_ok=True)
    return Response(status_code=204)

